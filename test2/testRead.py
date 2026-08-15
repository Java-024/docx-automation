import docx

doc = docx.Document("document.docx")

listKey = [{"symbol":"#","mission":"start"},
           {"symbol":"+","mission":"value+"},
           {"symbol":"-","mission":"value-"},
           {"symbol":"*","mission":"replace"},
           {"symbol":"=","mission":"keep"},
           {"symbol":"/","mission":"purge"},
           {"symbol":"&","mission":"merge"},]
runsList = []

def readFile(doc):
    readResult = []
    numberParagraph = 0
    for paragraph in doc.paragraphs:
        numberRun = 0
        for run in paragraph.runs:
            text = run.text
            is_bold = run.bold
            is_italic = run.italic
            readResult = readResult + [{
                "text":text,
                "is_bold":is_bold,
                "is_italic":is_italic,
                "numParagraph":numberParagraph,
                "numRun":numberRun
            }]
            print(f"Text: {text}")
            numberRun += 1
        numberParagraph += 1
    return readResult

def searchKeyAndReplace(runs, listKey, func):
    numParagraph, numRun = 0, 0
    listSymbolKey = ""
    startReplaceSymbol = ""
    symbolKey = ""
    startFlag = False
    mergeSymbol = ""

    for indexListKey in range(len(listKey)):
        if listKey[indexListKey].get("mission") == "start":
            startReplaceSymbol = listKey[indexListKey].get("symbol")
        if listKey[indexListKey].get("mission") == "merge":
            mergeSymbol = listKey[indexListKey].get("symbol")
        listSymbolKey = listSymbolKey + str(listKey[indexListKey].get("symbol"))

    for indexRun in range(len(runs)):
        string = runs[indexRun].get("text")
        for numSymbol in range(len(string)):
            if string[numSymbol] == startReplaceSymbol or string[numSymbol] == mergeSymbol:
                startFlag = True

            if symbolKey != "" and startFlag == False:
                symbolKey = ""

            if (startFlag == True) and (string[numSymbol] in listSymbolKey) and (string[numSymbol] != startReplaceSymbol):
                if symbolKey == "":
                    symbolKey = str(string[numSymbol])
                if symbolKey == string[numSymbol]:
                    pass
                print(symbolKey)



def replaceByKey():
    pass

searchKeyAndReplace(runs=readFile(doc), listKey=listKey, func=replaceByKey)

""""
for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        text = run.text
        is_bold = run.bold
        is_italic = run.italic
        print(f"Текст: {text} , Жирный: {is_bold} , Курсив: {is_italic}.")
        runsList = runsList + [text, paragraph, is_bold, is_italic]
"""

print(readFile(doc=doc))