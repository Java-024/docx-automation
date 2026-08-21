# импортирование модулей
from re import search

import docx
import os
import time

path = f"{os.getcwd()}\\original"
item = os.listdir(path)

print(item)

# чтение
doc = docx.Document(f"{path}\\{item[1]}")

print(len(doc.paragraphs))
print(len(doc.paragraphs[0].runs))

def search(string):
    first = -1
    endSimvol = -1
    if len(string) >2:
        if (string.find("[") != -1) and (string.find("]") != -1):
            for simvol in range(len(string)):
                if (string[simvol] == "[") and (first == -1) and (endSimvol == -1):
                    first = simvol
                if (string[simvol] == "]") and (first != -1) and (endSimvol == -1):
                    endSimvol = simvol
            if (first != -1) and (endSimvol != -1):
                return string[first:(endSimvol+1)]
            else: return None
    return False


def search1(string):
    result = []
    i = 0
    n = len(string)
    while i < n:
        if string[i] == '[':
            start = i
            balance = 0
            j = i
            found = False

            while j < n:
                if string[j] == '[':
                    balance += 1
                elif string[j] == ']':
                    balance -= 1
                    if balance == 0:
                        result.append(string[start:j + 1])
                        i = j + 1
                        found = True
                        break
                j += 1
            if not found: i += 1
        else: i += 1

    if not result:
        return False
    return result

def transform_list(input_list):
    """
    Преобразует список строк со специальными структурами в список списков.
    """
    result = []

    for item in input_list:
        # Убираем квадратные скобки
        content = item[1:-1].strip()

        # Разбиваем по запятой
        parts = content.split(',')

        # Очищаем каждую часть от кавычек и пробелов
        cleaned_parts = []
        for part in parts:
            # Убираем все кавычки и пробелы с начала и конца
            part = part.strip()
            # Заменяем все виды кавычек на пустоту с начала и конца
            while part and part[0] in ['"', '“', '”', "'"]:
                part = part[1:]
            while part and part[-1] in ['"', '“', '”', "'"]:
                part = part[:-1]
            part = part.strip()
            cleaned_parts.append(part)

        # Формируем результат
        if len(cleaned_parts) >= 2:
            result.append([cleaned_parts[0], cleaned_parts[1], item])
        else:
            result.append([item, '', item])

    return result


for i0 in range(len(doc.paragraphs)):
    time.sleep(0)
    #for i1 in range(len(doc.paragraphs[i0].runs)):
    vivod = search1(str(doc.paragraphs[i0].text))
    if isinstance(vivod, list):
        if vivod != []:
            print(f"{i0 + 1}. {transform_list(search1(str(doc.paragraphs[i0].text)))} ", end="")
    else:
        print(f"{i0 + 1}. {search1(str(doc.paragraphs[i0].text))} ", end="")
    print(doc.paragraphs[i0].text)
        #runs[i1].text)
        #pass

input()

# l = [['серияПаспортОтец', 'rd', '[“серияПаспортОтец”, “rd”]'], ['номерПаспортОтец', 'rd', '[“номерПаспортОтец”, “rd”]'], ['выданПаспортОтец', 'rd', '[“выданПаспортОтец”, “rd”]']]